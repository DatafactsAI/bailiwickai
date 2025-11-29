"""
Unit tests for PlaceholderExtractor class.
Tests placeholder extraction from Word documents and signature calculation.
"""

import pytest
import tempfile
import os
from pathlib import Path
from docx import Document

# Add parent directory to path for imports
import sys
sys.path.insert(0, str(Path(__file__).parent.parent))

from document_engine.placeholder_extractor import PlaceholderExtractor


@pytest.fixture
def extractor():
    """Create a PlaceholderExtractor instance."""
    return PlaceholderExtractor()


@pytest.fixture
def sample_docx_with_placeholders():
    """Create a sample .docx file with placeholders for testing."""
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        doc = Document()
        
        # Add paragraphs with placeholders
        doc.add_paragraph("Dear <client_name>,")
        doc.add_paragraph("This letter is dated <date> and prepared by <advisor_name>.")
        doc.add_paragraph("Your current super balance is <super_balance>.")
        doc.add_paragraph("Your gross salary is <gross_salary> per annum.")
        
        # Add a table with placeholders
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Client 1"
        table.cell(0, 1).text = "<client1_name>"
        table.cell(1, 0).text = "Client 2"
        table.cell(1, 1).text = "<client2_name>"
        
        doc.save(tmp.name)
        yield tmp.name
        
    # Cleanup
    os.unlink(tmp.name)


@pytest.fixture
def sample_docx_no_placeholders():
    """Create a sample .docx file without placeholders."""
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        doc = Document()
        doc.add_paragraph("This is a regular document.")
        doc.add_paragraph("It has no placeholders at all.")
        doc.save(tmp.name)
        yield tmp.name
        
    os.unlink(tmp.name)


@pytest.fixture
def sample_docx_duplicate_placeholders():
    """Create a sample .docx file with duplicate placeholders."""
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        doc = Document()
        doc.add_paragraph("Hello <client_name>!")
        doc.add_paragraph("Again, <client_name> is the name.")
        doc.add_paragraph("The date is <date> and again <date>.")
        doc.save(tmp.name)
        yield tmp.name
        
    os.unlink(tmp.name)


class TestPlaceholderExtraction:
    """Tests for placeholder extraction functionality."""
    
    def test_extract_placeholders_basic(self, extractor, sample_docx_with_placeholders):
        """Test basic placeholder extraction from paragraphs."""
        placeholders = extractor.extract_placeholders(sample_docx_with_placeholders)
        
        assert isinstance(placeholders, list)
        assert len(placeholders) > 0
        assert "<client_name>" in placeholders
        assert "<date>" in placeholders
        assert "<advisor_name>" in placeholders
    
    def test_extract_placeholders_from_tables(self, extractor, sample_docx_with_placeholders):
        """Test placeholder extraction from tables."""
        placeholders = extractor.extract_placeholders(sample_docx_with_placeholders)
        
        assert "<client1_name>" in placeholders
        assert "<client2_name>" in placeholders
    
    def test_extract_no_placeholders(self, extractor, sample_docx_no_placeholders):
        """Test extraction from document with no placeholders."""
        placeholders = extractor.extract_placeholders(sample_docx_no_placeholders)
        
        assert isinstance(placeholders, list)
        assert len(placeholders) == 0
    
    def test_extract_deduplicates_placeholders(self, extractor, sample_docx_duplicate_placeholders):
        """Test that duplicate placeholders are deduplicated."""
        placeholders = extractor.extract_placeholders(sample_docx_duplicate_placeholders)
        
        # Should only have unique placeholders
        assert placeholders.count("<client_name>") == 1
        assert placeholders.count("<date>") == 1
        assert len(placeholders) == 2
    
    def test_extract_sorted_alphabetically(self, extractor, sample_docx_with_placeholders):
        """Test that placeholders are sorted alphabetically."""
        placeholders = extractor.extract_placeholders(sample_docx_with_placeholders)
        
        # Should be sorted
        assert placeholders == sorted(placeholders)
    
    def test_extract_file_not_found(self, extractor):
        """Test error handling for non-existent file."""
        with pytest.raises(FileNotFoundError):
            extractor.extract_placeholders("/non/existent/file.docx")
    
    def test_extract_invalid_file_type(self, extractor):
        """Test error handling for non-.docx file."""
        with tempfile.NamedTemporaryFile(suffix='.txt', delete=False) as tmp:
            tmp.write(b"This is a text file")
            tmp.flush()
            
            with pytest.raises(ValueError):
                extractor.extract_placeholders(tmp.name)
            
            os.unlink(tmp.name)


class TestSignatureCalculation:
    """Tests for placeholder signature calculation."""
    
    def test_calculate_signature_basic(self, extractor):
        """Test basic signature calculation."""
        placeholders = ["<a>", "<b>", "<c>"]
        signature = extractor.calculate_signature(placeholders)
        
        assert isinstance(signature, str)
        assert len(signature) == 32  # MD5 hex length
    
    def test_calculate_signature_order_independent(self, extractor):
        """Test that signature is order-independent (sorted internally)."""
        sig1 = extractor.calculate_signature(["<a>", "<b>", "<c>"])
        sig2 = extractor.calculate_signature(["<c>", "<a>", "<b>"])
        
        assert sig1 == sig2
    
    def test_calculate_signature_different_for_different_inputs(self, extractor):
        """Test that different placeholder sets produce different signatures."""
        sig1 = extractor.calculate_signature(["<a>", "<b>"])
        sig2 = extractor.calculate_signature(["<a>", "<c>"])
        
        assert sig1 != sig2
    
    def test_calculate_signature_consistent(self, extractor):
        """Test that signature is consistent for same input."""
        placeholders = ["<client_name>", "<date>", "<advisor>"]
        
        sig1 = extractor.calculate_signature(placeholders)
        sig2 = extractor.calculate_signature(placeholders)
        
        assert sig1 == sig2
    
    def test_calculate_signature_empty_list(self, extractor):
        """Test signature calculation for empty list."""
        signature = extractor.calculate_signature([])
        
        assert isinstance(signature, str)
        assert len(signature) == 32


class TestDocumentStats:
    """Tests for document statistics calculation."""
    
    def test_get_document_stats(self, extractor, sample_docx_with_placeholders):
        """Test document statistics retrieval."""
        stats = extractor.get_document_stats(sample_docx_with_placeholders)
        
        assert isinstance(stats, dict)
        assert "word_count" in stats
        assert "paragraph_count" in stats
        assert "table_count" in stats
        
        assert stats["word_count"] > 0
        assert stats["paragraph_count"] > 0
        assert stats["table_count"] >= 1  # We added a table
    
    def test_get_document_stats_file_not_found(self, extractor):
        """Test stats for non-existent file."""
        with pytest.raises(FileNotFoundError):
            extractor.get_document_stats("/non/existent/file.docx")


if __name__ == "__main__":
    pytest.main([__file__, "-v"])


