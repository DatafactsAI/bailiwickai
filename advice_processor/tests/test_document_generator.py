"""
Unit tests for DocumentGenerator class.
Tests placeholder replacement and document generation.
"""

import pytest
import tempfile
import os
from pathlib import Path
from docx import Document

# Add parent directory to path for imports
import sys
sys.path.insert(0, str(Path(__file__).parent.parent))

from document_engine.document_generator import DocumentGenerator


@pytest.fixture
def generator():
    """Create a DocumentGenerator instance."""
    return DocumentGenerator()


@pytest.fixture
def simple_template():
    """Create a simple template with placeholders."""
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        doc = Document()
        doc.add_paragraph("Dear <client_name>,")
        doc.add_paragraph("This letter is dated <date>.")
        doc.add_paragraph("Sincerely, <advisor_name>")
        doc.save(tmp.name)
        yield tmp.name
    os.unlink(tmp.name)


@pytest.fixture
def table_template():
    """Create a template with table containing placeholders."""
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        doc = Document()
        doc.add_paragraph("Client Summary")
        
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "<client_name>"
        table.cell(1, 0).text = "Super Balance"
        table.cell(1, 1).text = "<super_balance>"
        table.cell(2, 0).text = "Advisor"
        table.cell(2, 1).text = "<advisor_name>"
        
        doc.save(tmp.name)
        yield tmp.name
    os.unlink(tmp.name)


@pytest.fixture
def formatted_template():
    """Create a template with formatted text."""
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        doc = Document()
        
        para = doc.add_paragraph()
        para.add_run("Dear ").bold = False
        para.add_run("<client_name>").bold = True
        para.add_run(",")
        
        doc.save(tmp.name)
        yield tmp.name
    os.unlink(tmp.name)


@pytest.fixture
def sample_values():
    """Sample resolved values for testing."""
    return {
        "<client_name>": "John Smith",
        "<date>": "15/01/2025",
        "<advisor_name>": "Andrea Paynter",
        "<super_balance>": "$320,000.00"
    }


class TestDocumentGeneration:
    """Tests for document generation functionality."""
    
    def test_generate_simple_document(self, generator, simple_template, sample_values):
        """Test basic document generation."""
        result = generator.generate_document(simple_template, sample_values)
        
        assert isinstance(result, bytes)
        assert len(result) > 0
    
    def test_generate_replaces_placeholders(self, generator, simple_template, sample_values):
        """Test that placeholders are actually replaced."""
        result_bytes = generator.generate_document(simple_template, sample_values)
        
        # Save to temp file and read back
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp.write(result_bytes)
            tmp.flush()
            
            doc = Document(tmp.name)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            
            # Placeholders should be replaced
            assert "<client_name>" not in full_text
            assert "<date>" not in full_text
            assert "<advisor_name>" not in full_text
            
            # Values should be present
            assert "John Smith" in full_text
            assert "15/01/2025" in full_text
            assert "Andrea Paynter" in full_text
            
            os.unlink(tmp.name)
    
    def test_generate_replaces_in_tables(self, generator, table_template, sample_values):
        """Test that placeholders in tables are replaced."""
        result_bytes = generator.generate_document(table_template, sample_values)
        
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp.write(result_bytes)
            tmp.flush()
            
            doc = Document(tmp.name)
            
            # Get text from all tables
            table_text = ""
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        table_text += cell.text + " "
            
            # Check replacements
            assert "John Smith" in table_text
            assert "$320,000.00" in table_text
            assert "Andrea Paynter" in table_text
            
            # No placeholders should remain
            assert "<client_name>" not in table_text
            assert "<super_balance>" not in table_text
            
            os.unlink(tmp.name)
    
    def test_generate_with_empty_values(self, generator, simple_template):
        """Test generation with empty string values."""
        values = {
            "<client_name>": "",
            "<date>": "",
            "<advisor_name>": ""
        }
        
        result_bytes = generator.generate_document(simple_template, values)
        
        assert isinstance(result_bytes, bytes)
        assert len(result_bytes) > 0
        
        # Verify document can be opened
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp.write(result_bytes)
            tmp.flush()
            doc = Document(tmp.name)
            assert len(doc.paragraphs) > 0
            os.unlink(tmp.name)
    
    def test_generate_with_partial_values(self, generator, simple_template):
        """Test generation with only some placeholders mapped."""
        values = {
            "<client_name>": "John Smith"
            # <date> and <advisor_name> not provided
        }
        
        result_bytes = generator.generate_document(simple_template, values)
        
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp.write(result_bytes)
            tmp.flush()
            
            doc = Document(tmp.name)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            
            # Replaced placeholder should have value
            assert "John Smith" in full_text
            
            # Unreplaced placeholders should remain
            assert "<date>" in full_text
            assert "<advisor_name>" in full_text
            
            os.unlink(tmp.name)
    
    def test_generate_file_not_found(self, generator, sample_values):
        """Test error handling for non-existent template."""
        with pytest.raises(FileNotFoundError):
            generator.generate_document("/non/existent/file.docx", sample_values)


class TestPreviewReplacements:
    """Tests for preview functionality."""
    
    def test_preview_simple(self, generator, simple_template, sample_values):
        """Test preview of replacements."""
        preview = generator.preview_replacements(simple_template, sample_values)
        
        assert isinstance(preview, dict)
        assert "body" in preview
        assert "tables" in preview
        assert "headers" in preview
        assert "footers" in preview
    
    def test_preview_finds_replacements(self, generator, simple_template, sample_values):
        """Test that preview finds all replacements."""
        preview = generator.preview_replacements(simple_template, sample_values)
        
        # Should find body replacements
        assert len(preview["body"]) > 0
        
        # Each replacement should have placeholder and value
        for replacement in preview["body"]:
            assert "placeholder" in replacement
            assert "value" in replacement
    
    def test_preview_table_replacements(self, generator, table_template, sample_values):
        """Test preview finds table replacements."""
        preview = generator.preview_replacements(table_template, sample_values)
        
        assert len(preview["tables"]) > 0


class TestEdgeCases:
    """Tests for edge cases and special scenarios."""
    
    def test_placeholder_with_special_chars(self, generator):
        """Test placeholders with underscores and hyphens."""
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            doc = Document()
            doc.add_paragraph("Value: <my_special-placeholder>")
            doc.save(tmp.name)
            
            values = {"<my_special-placeholder>": "replaced_value"}
            result_bytes = generator.generate_document(tmp.name, values)
            
            # Read and verify
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp2:
                tmp2.write(result_bytes)
                tmp2.flush()
                doc2 = Document(tmp2.name)
                
                assert "replaced_value" in doc2.paragraphs[0].text
                os.unlink(tmp2.name)
            
            os.unlink(tmp.name)
    
    def test_multiple_same_placeholder(self, generator):
        """Test replacing multiple occurrences of the same placeholder."""
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            doc = Document()
            doc.add_paragraph("Hello <name>! Nice to meet you, <name>.")
            doc.save(tmp.name)
            
            values = {"<name>": "John"}
            result_bytes = generator.generate_document(tmp.name, values)
            
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp2:
                tmp2.write(result_bytes)
                tmp2.flush()
                doc2 = Document(tmp2.name)
                
                text = doc2.paragraphs[0].text
                assert text.count("John") == 2
                assert "<name>" not in text
                os.unlink(tmp2.name)
            
            os.unlink(tmp.name)
    
    def test_empty_document(self, generator):
        """Test generating from an empty document."""
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            doc = Document()
            doc.save(tmp.name)
            
            result_bytes = generator.generate_document(tmp.name, {})
            
            assert isinstance(result_bytes, bytes)
            assert len(result_bytes) > 0
            
            os.unlink(tmp.name)


if __name__ == "__main__":
    pytest.main([__file__, "-v"])


