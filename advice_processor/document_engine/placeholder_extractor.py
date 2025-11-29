"""
Placeholder Extractor - Extracts template placeholders from Word documents.
Finds all text patterns matching <placeholder_name> format in paragraphs, tables, headers, and footers.
"""

import re
import hashlib
import logging
from pathlib import Path
from typing import List, Set
from docx import Document
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)


class PlaceholderExtractor:
    """
    Extracts placeholders from Word documents (.docx).
    Placeholders are identified by the pattern <placeholder_name>.
    """
    
    # Regex pattern to match placeholders like <client_name>, <DATE>, <custom-field>
    PLACEHOLDER_PATTERN = re.compile(r'<([a-zA-Z][a-zA-Z0-9_-]*)>')
    
    def __init__(self):
        pass
    
    def extract_placeholders(self, docx_path: str) -> List[str]:
        """
        Extract all unique placeholders from a Word document.
        
        Args:
            docx_path: Path to the .docx file
            
        Returns:
            List of unique placeholder strings (e.g., ["<client_name>", "<date>"])
        """
        path = Path(docx_path)
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {docx_path}")
        
        if not path.suffix.lower() == '.docx':
            raise ValueError(f"Invalid file type: {path.suffix}. Expected .docx")
        
        try:
            doc = Document(docx_path)
        except Exception as e:
            logger.error(f"Failed to open document: {str(e)}")
            raise ValueError(f"Failed to open document: {str(e)}")
        
        placeholders: Set[str] = set()
        
        # Extract from main document body paragraphs
        for paragraph in doc.paragraphs:
            found = self._extract_from_text(paragraph.text)
            placeholders.update(found)
        
        # Extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        found = self._extract_from_text(paragraph.text)
                        placeholders.update(found)
        
        # Extract from headers and footers
        for section in doc.sections:
            # Header
            if section.header:
                for paragraph in section.header.paragraphs:
                    found = self._extract_from_text(paragraph.text)
                    placeholders.update(found)
                # Tables in header
                for table in section.header.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                found = self._extract_from_text(paragraph.text)
                                placeholders.update(found)
            
            # Footer
            if section.footer:
                for paragraph in section.footer.paragraphs:
                    found = self._extract_from_text(paragraph.text)
                    placeholders.update(found)
                # Tables in footer
                for table in section.footer.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                found = self._extract_from_text(paragraph.text)
                                placeholders.update(found)
        
        # Sort alphabetically for consistent ordering
        sorted_placeholders = sorted(list(placeholders))
        
        logger.info(f"Extracted {len(sorted_placeholders)} unique placeholders from {path.name}")
        return sorted_placeholders
    
    def _extract_from_text(self, text: str) -> List[str]:
        """
        Extract placeholders from a text string.
        
        Args:
            text: Text to search for placeholders
            
        Returns:
            List of placeholder strings found (including angle brackets)
        """
        if not text:
            return []
        
        matches = self.PLACEHOLDER_PATTERN.findall(text)
        # Return full placeholder format with angle brackets
        return [f"<{match}>" for match in matches]
    
    def calculate_signature(self, placeholders: List[str]) -> str:
        """
        Calculate a signature hash for a list of placeholders.
        Used to detect if a template's placeholders have changed.
        
        Args:
            placeholders: List of placeholder strings
            
        Returns:
            MD5 hash string of the sorted placeholders
        """
        sorted_placeholders = sorted(placeholders)
        signature_string = "|".join(sorted_placeholders)
        return hashlib.md5(signature_string.encode()).hexdigest()
    
    def get_document_stats(self, docx_path: str) -> dict:
        """
        Get basic statistics about a Word document.
        
        Args:
            docx_path: Path to the .docx file
            
        Returns:
            Dictionary with word_count, paragraph_count, table_count
        """
        path = Path(docx_path)
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {docx_path}")
        
        try:
            doc = Document(docx_path)
        except Exception as e:
            logger.error(f"Failed to open document: {str(e)}")
            return {"word_count": 0, "paragraph_count": 0, "table_count": 0}
        
        word_count = 0
        paragraph_count = len(doc.paragraphs)
        table_count = len(doc.tables)
        
        # Count words in paragraphs
        for paragraph in doc.paragraphs:
            words = paragraph.text.split()
            word_count += len(words)
        
        # Count words in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        words = paragraph.text.split()
                        word_count += len(words)
        
        return {
            "word_count": word_count,
            "paragraph_count": paragraph_count,
            "table_count": table_count
        }


